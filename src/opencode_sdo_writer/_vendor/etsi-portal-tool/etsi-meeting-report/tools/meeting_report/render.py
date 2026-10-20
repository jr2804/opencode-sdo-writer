"""Render the agenda body into the meeting's own document, and compare a rendered document with a reference.

The head of the document (cover, title block, schedule, IPR and Code of Conduct boilerplate) is kept
from the meeting's own document. Everything after the anchor paragraph is replaced by the body
template, rendered with the meeting context.
"""
from __future__ import annotations

import difflib
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docxtpl import DocxTemplate


def _truncate_after(doc, anchor_text: str) -> None:
    """Remove everything after the paragraph that starts with anchor_text (section properties kept)."""
    body = doc.element.body
    anchor = next((p._p for p in doc.paragraphs if p.text.startswith(anchor_text)), None)
    if anchor is None:
        raise SystemExit(f"anchor paragraph not found in the document: {anchor_text!r}")
    sect = body.find(qn("w:sectPr"))
    el = anchor.getnext()
    while el is not None and el is not sect:
        nxt = el.getnext()
        body.remove(el)
        el = nxt


def render_agenda(base: Path, body_template: Path, context: dict, anchor: str, out: Path) -> None:
    """Head from ``base`` up to ``anchor``; body from ``body_template`` rendered with ``context``."""
    scratch = out.with_name(out.stem + ".body.docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    tpl = DocxTemplate(body_template)
    tpl.render(context)
    tpl.save(scratch)

    doc = Document(base)
    _truncate_after(doc, anchor)
    sect = doc.element.body.find(qn("w:sectPr"))
    rendered = Document(scratch)
    for el in list(rendered.element.body.iterchildren()):
        if el.tag == qn("w:sectPr"):
            continue
        sect.addprevious(deepcopy(el))
    doc.save(out)
    scratch.unlink()


def body_lines(path: Path, start: str = "Approval of agenda") -> list[str]:
    """The agenda body (from the start heading to the end) as comparable lines."""
    doc = Document(path)
    out: list[str] = []
    started = False
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc)
            if not started:
                if p.text.strip() == start:
                    started = True
                else:
                    continue
            out.append(f"P[{p.style.name}] {p.text}")
        elif el.tag == qn("w:tbl") and started:
            for row in Table(el, doc).rows:
                out.append("T| " + " | ".join(c.text for c in row.cells))
    return out


def compare(reference: Path, rendered: Path) -> tuple[list[str], int]:
    a, b = body_lines(reference), body_lines(rendered)
    diff = list(difflib.unified_diff(a, b, "reference", "rendered", lineterm="", n=1))
    return diff, sum(1 for line in diff if line[:1] in "+-" and not line.startswith(("+++", "---")))
